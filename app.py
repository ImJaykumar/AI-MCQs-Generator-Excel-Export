import json
import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import streamlit as st

st.set_page_config(page_title="JSON to Excel Converter", page_icon="📊", layout="centered")

st.title("📊 JSON to Excel Converter")
st.write("Upload your Hindi MCQ JSON file to convert it into a beautifully formatted Excel sheet.")

uploaded_file = st.file_uploader("Choose a JSON file", type=["json"])

def convert_json_to_excel(json_data):
    formatted_rows = []

    for idx, item in enumerate(json_data, start=1):
        options = item.get('options', [])
        
        opt_a = options[0] if len(options) > 0 else ''
        opt_b = options[1] if len(options) > 1 else ''
        opt_c = options[2] if len(options) > 2 else ''
        opt_d = options[3] if len(options) > 3 else ''

        row = {
            'Sr. No.': idx,
            'Question': item.get('question', ''),
            'Option A': opt_a,
            'Option B': opt_b,
            'Option C': opt_c,
            'Option D': opt_d,
            'Answer': item.get('correct_answer', ''),
            'Explanation': item.get('explanation', '')
        }
        formatted_rows.append(row)

    df = pd.DataFrame(formatted_rows)

    # Save DataFrame to an in-memory Excel file using openpyxl
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='MCQs')
        
        # Access the openpyxl worksheet to apply auto-width and styling
        worksheet = writer.sheets['MCQs']
        
        # Header Style
        header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid') # Dark Blue Header
        
        for col_num in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # Auto-adjust column widths based on maximum length of content
        for col in worksheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            
            for cell in col:
                val = str(cell.value or '')
                # Handle line breaks gracefully when measuring length
                lines = val.split('\n')
                line_max = max(len(l) for l in lines) if lines else 0
                if line_max > max_len:
                    max_len = line_max
            
            # Set width with extra padding, capped at a maximum width of 60 for long questions/explanations
            adjusted_width = min(max(max_len + 4, 12), 60)
            worksheet.column_dimensions[col_letter].width = adjusted_width

    output.seek(0)
    return output

if uploaded_file is not None:
    try:
        json_data = json.load(uploaded_file)
        
        # Process and convert
        excel_data = convert_json_to_excel(json_data)
        
        st.success("JSON processed and Excel sheet generated successfully!")
        
        # Download Button
        file_name = uploaded_file.name.rsplit('.', 1)[0] + "_formatted.xlsx"
        st.download_button(
            label="📥 Download Excel File",
            data=excel_data,
            file_name=file_name,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    except Exception as e:
        st.error(f"Error processing file: {e}")